"""
Export the structured report (see src/report/build.py) to a Word (.docx)
document, so the investigation report is a shareable file, not only
on-screen text.
"""
import json
import os
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.report.build import build_report

GROUNDED_COLOR = RGBColor(0x1A, 0x7F, 0x37)
UNGROUNDED_COLOR = RGBColor(0xB3, 0x26, 0x1E)


def _add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, "") else "-"
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table


def export_docx(report, out_path="output/report.docx", graph_path=None):
    graph_path = graph_path or report.get("graph_path", "graph.png")
    doc = Document()

    title = doc.add_heading("Digital Forensic Investigation Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    meta.add_run(f"Case: {report.get('case_id', 'unknown')}").bold = True
    meta.add_run(f"    Generated: {report.get('generated_at', '')}")

    # --- 1. Executive summary ------------------------------------------------
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(report.get("executive_summary", "No summary produced."))

    # --- 2. Timeline -----------------------------------------------------------
    doc.add_heading("2. Timeline", level=1)
    timeline = report.get("timeline", [])
    if timeline:
        rows = [
            [t["timestamp"], t["artifact_id"], "SUSPICIOUS" if t["suspicious"] else "",
             t["actor"], t["event_type"], t["object"], t["command"], t["dst_ip"],
             "; ".join(t["techniques"])]
            for t in timeline
        ]
        _add_table(doc, ["Timestamp", "Artifact ID", "Flag", "Actor", "Event Type",
                          "Object", "Command", "Dst IP", "MITRE"], rows)
    else:
        doc.add_paragraph("No events available.")

    # --- 3. Techniques detected -------------------------------------------------
    doc.add_heading("3. Techniques Detected", level=1)
    techs = report.get("techniques_detected", [])
    if techs:
        rows = [
            [t["technique_id"], t["technique_name"], t["count"], ", ".join(t["artifact_ids"])]
            for t in techs
        ]
        _add_table(doc, ["Technique ID", "Technique Name", "Event Count", "Artifact IDs"], rows)
    else:
        doc.add_paragraph("No MITRE ATT&CK techniques were matched against this event set.")

    # --- 4. Attack narrative -----------------------------------------------------
    doc.add_heading("4. Attack Narrative (Evidence-Grounded)", level=1)
    if graph_path and os.path.exists(graph_path):
        try:
            doc.add_picture(graph_path, width=Inches(6))
        except Exception:
            pass

    for s in report.get("narrative", []):
        h = doc.add_heading(level=2)
        run = h.add_run(s["stage"])
        mark = h.add_run(f"  [{'OK' if s['grounded'] else 'UNSUPPORTED'}]")
        mark.font.color.rgb = GROUNDED_COLOR if s["grounded"] else UNGROUNDED_COLOR
        mark.bold = True

        doc.add_paragraph(s["claim"])
        ids = ", ".join(s["artifact_ids"]) or "none"
        p = doc.add_paragraph()
        p.add_run("Evidence: ").bold = True
        p.add_run(ids)
        p2 = doc.add_paragraph()
        p2.add_run("Confidence (model-reported): ").bold = True
        p2.add_run(str(s["confidence"]))
        if s["missing_ids"]:
            warn = doc.add_paragraph()
            warn.add_run(f"WARNING: cited non-existent evidence: {', '.join(s['missing_ids'])}").font.color.rgb = UNGROUNDED_COLOR

    if report.get("unsupported_claims"):
        doc.add_heading("Flagged unsupported claims", level=2)
        doc.add_paragraph("The following claims were NOT backed by real evidence and should not be trusted:")
        for c in report["unsupported_claims"]:
            doc.add_paragraph(c, style="List Bullet")

    # --- 5. Grounding score -------------------------------------------------------
    doc.add_heading("5. Grounding Score", level=1)
    g = report.get("grounding", {})
    p = doc.add_paragraph()
    p.add_run(f"{g.get('score', 0)}").bold = True
    p.add_run(f"  ({g.get('grounded_claims', 0)}/{g.get('total_claims', 0)} "
               f"claims backed by cited evidence)")

    # --- 6. Evidence appendix ------------------------------------------------------
    doc.add_heading("6. Evidence Appendix", level=1)
    appendix = report.get("evidence_appendix", [])
    if appendix:
        rows = [
            [a["artifact_id"], a["timestamp"], a["source"], a["actor"], a["event_type"],
             a["object"], a["command"], a["src_ip"], a["dst_ip"]]
            for a in appendix
        ]
        _add_table(doc, ["Artifact ID", "Timestamp", "Source", "Actor", "Event Type",
                          "Object", "Command", "Src IP", "Dst IP"], rows)
        doc.add_paragraph("")
        doc.add_heading("Raw log lines", level=3)
        for a in appendix:
            doc.add_paragraph(f"{a['artifact_id']}: {a['raw']}", style="List Bullet")
    else:
        doc.add_paragraph("No events recorded.")

    # --- Limitations -----------------------------------------------------------------
    doc.add_heading("Limitations", level=1)
    doc.add_paragraph(report.get("limitations", "None stated."))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


def run(verified_path, correlations_path, out_path="output/report.docx"):
    with open(verified_path) as f:
        verified = json.load(f)
    with open(correlations_path) as f:
        corr = json.load(f)
    report = build_report(verified, corr.get("events", []))
    return export_docx(report, out_path=out_path)


if __name__ == "__main__":
    verified_path = sys.argv[1] if len(sys.argv) > 1 else "output/verified.json"
    correlations_path = sys.argv[2] if len(sys.argv) > 2 else "data/samples/correlations.json"
    out = run(verified_path, correlations_path)
    print(f"--- saved {out} ---")